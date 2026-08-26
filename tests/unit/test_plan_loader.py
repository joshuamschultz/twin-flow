"""COMP-018 PlanLoader — unit tests (T-030, red phase).

Contract under test (D-054, SDD Layer 3):
  - `load_plan(path, registry)` reads ONE tabular sheet per file: the first/only
    worksheet of an `.xlsx`, or a `.csv` with identical columns. Column contract
    only — no transformation logic, no routing/timing math.
  - Required columns: work_order_id, part, qty, start_date, due_date.
  - Optional initial-WIP columns: initial_wip_location, initial_wip_qty,
    initial_wip_remaining_time — a bundle already on the floor at a named
    location with remaining qty and time.
  - Returns list[WorkOrder]; a `.xlsx` and a `.csv` with the same rows load
    IDENTICALLY.
  - Validated against a PartTypeRegistry: an unknown part is rejected naming
    the row number; a missing required column is rejected.
  - No spreadsheet formula is ever evaluated to a computed result.

API committed by these tests (implementer conforms, names kept boring):
  - `WorkOrder` — a dataclass with fields work_order_id: str, part: str,
    qty: int, start_date: str, due_date: str, initial_wip_location: str | None
    = None, initial_wip_qty: int | None = None,
    initial_wip_remaining_time: float | None = None. Equality is structural
    (dataclass default `__eq__`).
  - Every rejection (unknown part, missing required column, formula cell,
    malformed value) raises `ValueError` (or a subclass of it — a domain
    exception like the codebase's `ExpressionError` pattern is fine, as long
    as it `isinstance`-matches `ValueError`, since that is the only name
    these tests import and the stub currently exports). The message names
    the row number where the failure applies to one row, using the literal
    substring "row N" (1-indexed with the header as row 1) so a message
    cannot accidentally satisfy the assertion by coincidence.
  - `load_plan(path: str, registry: PartTypeRegistry) -> list[WorkOrder]`.

Design choices this file locks in, and are flagged for the builder to confirm
rather than silently assumed:
  - start_date/due_date are read as-is (plain strings) with NO date parsing —
    D-054 says "column contract only, no transformation logic", and the SDD
    never specifies a date type. Fixtures therefore write plain ISO strings
    into every date cell (never an Excel date-formatted cell), so parity
    between .xlsx and .csv does not depend on a parsing decision this test
    file has no authority to make.
  - qty / initial_wip_qty / initial_wip_remaining_time DO need numeric
    coercion, because .csv is text-only and .xlsx numeric cells are already
    typed — without coercion the two formats could never load identically.
    This is column *typing*, not transformation logic.
  - A formula cell is REJECTED outright (a ValueError), not merely
    "read as its formula string and passed through". Passing a formula
    string through to a numeric field only defers the danger to whatever
    reads it next; failing closed here matches D-007.
  - A row where SOME but not ALL of the three initial_wip_* columns are
    populated is rejected — a half-specified WIP placement (e.g. a location
    with no remaining qty) is not a valid "no WIP" row and not a valid WIP
    row either.
  - A non-numeric qty value is rejected naming its row (fail closed on
    malformed input, D-007), not silently coerced or defaulted.

These are genuine open questions this file surfaces for the builder/
orchestrator to confirm — not smuggled requirements.
"""

from __future__ import annotations

import csv
from pathlib import Path

import openpyxl
import pytest

from factory_twin.plan.loader import WorkOrder, load_plan
from factory_twin.primitives.part import PartTypeRegistry

# ---------------------------------------------------------------------------
# Shared fixtures
# ---------------------------------------------------------------------------

REQUIRED_COLUMNS = ["work_order_id", "part", "qty", "start_date", "due_date"]
ALL_COLUMNS = REQUIRED_COLUMNS + [
    "initial_wip_location",
    "initial_wip_qty",
    "initial_wip_remaining_time",
]


def _registry() -> PartTypeRegistry:
    return PartTypeRegistry(
        {
            "blank": {"attributes": {"length": float}, "uom": "piece"},
            "wire": {"attributes": {"gauge": float}, "uom": "ft"},
        }
    )


def _write_csv(path: Path, header: list[str], rows: list[dict[str, object]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=header)
        writer.writeheader()
        for row in rows:
            writer.writerow({col: ("" if row.get(col) is None else row.get(col)) for col in header})


def _write_xlsx(path: Path, header: list[str], rows: list[dict[str, object]]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append([row.get(col) for col in header])
    workbook.save(path)


# Two ordinary orders, no WIP, plus one order with a fully-populated WIP row.
_IDENTICAL_PLAN_ROWS: list[dict[str, object]] = [
    {
        "work_order_id": "WO-1001",
        "part": "blank",
        "qty": 50,
        "start_date": "2026-09-01",
        "due_date": "2026-09-10",
    },
    {
        "work_order_id": "WO-1002",
        "part": "wire",
        "qty": 12,
        "start_date": "2026-09-02",
        "due_date": "2026-09-15",
        "initial_wip_location": "ht_oven",
        "initial_wip_qty": 8,
        "initial_wip_remaining_time": 4.5,
    },
]


# ---------------------------------------------------------------------------
# Criterion 1 — .xlsx and .csv load IDENTICALLY
# ---------------------------------------------------------------------------


def test_load_plan_xlsx_and_csv_with_same_rows_produce_equal_work_orders(
    tmp_path: Path,
) -> None:
    csv_path = tmp_path / "plan.csv"
    xlsx_path = tmp_path / "plan.xlsx"
    _write_csv(csv_path, ALL_COLUMNS, _IDENTICAL_PLAN_ROWS)
    _write_xlsx(xlsx_path, ALL_COLUMNS, _IDENTICAL_PLAN_ROWS)
    registry = _registry()

    csv_orders = load_plan(str(csv_path), registry)
    xlsx_orders = load_plan(str(xlsx_path), registry)

    assert csv_orders == xlsx_orders


def test_load_plan_csv_returns_expected_work_order_fields(tmp_path: Path) -> None:
    csv_path = tmp_path / "plan.csv"
    _write_csv(csv_path, ALL_COLUMNS, _IDENTICAL_PLAN_ROWS)
    registry = _registry()

    orders = load_plan(str(csv_path), registry)

    assert orders[0] == WorkOrder(
        work_order_id="WO-1001",
        part="blank",
        qty=50,
        start_date="2026-09-01",
        due_date="2026-09-10",
    )


def test_load_plan_reads_only_the_first_worksheet_of_a_multi_sheet_workbook(
    tmp_path: Path,
) -> None:
    """D-054: ONE table per file — the first/only worksheet. A second sheet with
    plan-shaped data must never be read; only the sheet openpyxl creates by
    default (and that this fixture populates) counts.
    """
    xlsx_path = tmp_path / "plan.xlsx"
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "orders"
    first.append(ALL_COLUMNS)
    first.append(
        [
            "WO-2001",
            "blank",
            5,
            "2026-09-01",
            "2026-09-05",
            None,
            None,
            None,
        ]
    )
    second = workbook.create_sheet("decoy")
    second.append(ALL_COLUMNS)
    second.append(
        [
            "WO-9999",
            "wire",
            999,
            "2026-01-01",
            "2026-01-02",
            None,
            None,
            None,
        ]
    )
    workbook.save(xlsx_path)
    registry = _registry()

    orders = load_plan(str(xlsx_path), registry)

    assert [order.work_order_id for order in orders] == ["WO-2001"]


# ---------------------------------------------------------------------------
# Criterion 2 — initial-WIP rows place a bundle at a named location
# ---------------------------------------------------------------------------


def test_load_plan_parses_initial_wip_location_qty_and_remaining_time(
    tmp_path: Path,
) -> None:
    csv_path = tmp_path / "plan.csv"
    _write_csv(csv_path, ALL_COLUMNS, _IDENTICAL_PLAN_ROWS)
    registry = _registry()

    orders = load_plan(str(csv_path), registry)
    wip_order = next(order for order in orders if order.work_order_id == "WO-1002")

    assert wip_order.initial_wip_location == "ht_oven"
    assert wip_order.initial_wip_qty == 8
    assert wip_order.initial_wip_remaining_time == 4.5


def test_load_plan_order_without_wip_columns_has_no_initial_wip(tmp_path: Path) -> None:
    csv_path = tmp_path / "plan.csv"
    _write_csv(csv_path, ALL_COLUMNS, _IDENTICAL_PLAN_ROWS)
    registry = _registry()

    orders = load_plan(str(csv_path), registry)
    plain_order = next(order for order in orders if order.work_order_id == "WO-1001")

    assert plain_order.initial_wip_location is None
    assert plain_order.initial_wip_qty is None
    assert plain_order.initial_wip_remaining_time is None


def test_load_plan_partially_populated_wip_columns_is_rejected(tmp_path: Path) -> None:
    """A location with no remaining qty (or vice versa) is not a valid row in
    either direction — reject rather than guess which fields the client meant.
    """
    csv_path = tmp_path / "plan.csv"
    rows = [
        {
            "work_order_id": "WO-3001",
            "part": "blank",
            "qty": 10,
            "start_date": "2026-09-01",
            "due_date": "2026-09-05",
            "initial_wip_location": "ht_oven",
            # initial_wip_qty and initial_wip_remaining_time left blank
        }
    ]
    _write_csv(csv_path, ALL_COLUMNS, rows)
    registry = _registry()

    with pytest.raises(ValueError, match="row 2"):
        load_plan(str(csv_path), registry)


# ---------------------------------------------------------------------------
# Criterion 3 — unknown part / missing column / formulas never evaluated
# ---------------------------------------------------------------------------


def test_load_plan_rejects_unknown_part_naming_the_row_number(tmp_path: Path) -> None:
    csv_path = tmp_path / "plan.csv"
    rows = [
        {
            "work_order_id": "WO-1001",
            "part": "blank",
            "qty": 5,
            "start_date": "2026-09-01",
            "due_date": "2026-09-05",
        },
        {
            "work_order_id": "WO-1002",
            "part": "unobtainium",  # not declared in the registry
            "qty": 5,
            "start_date": "2026-09-01",
            "due_date": "2026-09-05",
        },
    ]
    _write_csv(csv_path, ALL_COLUMNS, rows)
    registry = _registry()

    # header is row 1; the bad row is the second data row -> row 3.
    with pytest.raises(ValueError, match="row 3"):
        load_plan(str(csv_path), registry)


def test_load_plan_rejects_unknown_part_on_the_first_data_row(tmp_path: Path) -> None:
    """Off-by-one guard: the very first data row is row 2 (row 1 is the header),
    in both .csv and .xlsx.
    """
    csv_path = tmp_path / "plan.csv"
    rows = [
        {
            "work_order_id": "WO-1001",
            "part": "unobtainium",
            "qty": 5,
            "start_date": "2026-09-01",
            "due_date": "2026-09-05",
        }
    ]
    _write_csv(csv_path, ALL_COLUMNS, rows)
    registry = _registry()

    with pytest.raises(ValueError, match="row 2"):
        load_plan(str(csv_path), registry)


def test_load_plan_rejects_missing_required_column(tmp_path: Path) -> None:
    csv_path = tmp_path / "plan.csv"
    header_missing_due_date = [c for c in REQUIRED_COLUMNS if c != "due_date"]
    rows = [
        {
            "work_order_id": "WO-1001",
            "part": "blank",
            "qty": 5,
            "start_date": "2026-09-01",
        }
    ]
    _write_csv(csv_path, header_missing_due_date, rows)
    registry = _registry()

    with pytest.raises(ValueError, match="due_date"):
        load_plan(str(csv_path), registry)


def test_load_plan_never_evaluates_an_xlsx_formula_cell(tmp_path: Path) -> None:
    """A formula in the qty column must never be evaluated to its computed
    result: `=1+1` must never become the integer 2. This file's contract
    rejects the row outright rather than pass the formula text through.
    """
    xlsx_path = tmp_path / "plan.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(ALL_COLUMNS)
    sheet.append(
        [
            "WO-4001",
            "blank",
            "=1+1",  # formula string written directly into the qty cell
            "2026-09-01",
            "2026-09-05",
            None,
            None,
            None,
        ]
    )
    workbook.save(xlsx_path)
    registry = _registry()

    with pytest.raises(ValueError, match="row 2"):
        load_plan(str(xlsx_path), registry)


def test_load_plan_rejects_non_numeric_qty_naming_the_row(tmp_path: Path) -> None:
    csv_path = tmp_path / "plan.csv"
    rows = [
        {
            "work_order_id": "WO-5001",
            "part": "blank",
            "qty": "a-lot",
            "start_date": "2026-09-01",
            "due_date": "2026-09-05",
        }
    ]
    _write_csv(csv_path, ALL_COLUMNS, rows)
    registry = _registry()

    with pytest.raises(ValueError, match="row 2"):
        load_plan(str(csv_path), registry)


# ---------------------------------------------------------------------------
# Adversarial extras beyond the listed acceptance criteria
# ---------------------------------------------------------------------------


def test_load_plan_with_only_a_header_returns_an_empty_list(tmp_path: Path) -> None:
    csv_path = tmp_path / "plan.csv"
    _write_csv(csv_path, ALL_COLUMNS, [])
    registry = _registry()

    orders = load_plan(str(csv_path), registry)

    assert orders == []
